"""Convert reports/report_sections_4_5_6_8.md to a Word .docx file.

Uses python-docx with custom handling for:
- Headings (#, ##, ###)
- Tables (markdown pipe tables)
- Code blocks (```...```)
- Inline code (`...`)
- Bold (**...**) and italic (*...*)
- Bullet lists
- Plain paragraphs

Output: reports/report_sections_4_5_6_8.docx
"""
from __future__ import annotations
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "reports" / "report_sections_4_5_6_8.md"
DST = ROOT / "reports" / "report_sections_4_5_6_8.docx"


def add_inline_runs(paragraph, text: str):
    """Render inline markdown formatting (bold, italic, code) into runs."""
    # Tokenize: bold **x**, italic *x*, code `x`
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def add_table(doc, header: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(header):
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h.strip())
        run.bold = True
        run.font.size = Pt(10)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            p = cell.paragraphs[0]
            add_inline_runs(p, cell_text.strip())
            for run in p.runs:
                run.font.size = Pt(9.5)


def add_code_block(doc, code: str, lang: str = ""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # light gray background via shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F5F5F5")
    pPr.append(shd)


def parse_table_block(lines: list[str], i: int) -> tuple[list[str], list[list[str]], int]:
    """Parse a markdown table starting at lines[i] (header row).
    Returns (header, rows, next_index)."""
    def split_row(s: str) -> list[str]:
        s = s.strip()
        if s.startswith("|"):
            s = s[1:]
        if s.endswith("|"):
            s = s[:-1]
        return [c.strip() for c in s.split("|")]

    header = split_row(lines[i])
    # next line should be separator |---|---|
    sep = lines[i + 1] if i + 1 < len(lines) else ""
    if not re.match(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$", sep):
        return header, [], i + 1  # not actually a table
    rows = []
    j = i + 2
    while j < len(lines) and lines[j].strip().startswith("|"):
        rows.append(split_row(lines[j]))
        j += 1
    return header, rows, j


def convert():
    doc = Document()
    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    text = SRC.read_text(encoding="utf-8")
    lines = text.split("\n")

    i = 0
    in_code = False
    code_buf: list[str] = []
    code_lang = ""

    while i < len(lines):
        line = lines[i]

        # Code fence toggle
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line[3:].strip()
                code_buf = []
            else:
                add_code_block(doc, "\n".join(code_buf), code_lang)
                in_code = False
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2).strip()
            doc.add_heading(heading_text, level=min(level, 4))
            i += 1
            continue

        # Horizontal rule
        if re.match(r"^-{3,}\s*$", line):
            doc.add_paragraph()  # spacing
            i += 1
            continue

        # Tables
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            header, rows, next_i = parse_table_block(lines, i)
            add_table(doc, header, rows)
            doc.add_paragraph()
            i = next_i
            continue

        # Bullet list
        if re.match(r"^\s*[-*]\s+", line):
            bullet_text = re.sub(r"^\s*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, bullet_text)
            i += 1
            continue

        # Numbered list
        if re.match(r"^\s*\d+\.\s+", line):
            num_text = re.sub(r"^\s*\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, num_text)
            i += 1
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Plain paragraph: collect until blank line or special construct
        para_lines = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            if not nxt.strip():
                break
            if nxt.startswith("#") or nxt.startswith("```"):
                break
            if nxt.strip().startswith("|") or re.match(r"^\s*[-*]\s+", nxt) or re.match(r"^\s*\d+\.\s+", nxt):
                break
            if re.match(r"^-{3,}\s*$", nxt):
                break
            para_lines.append(nxt)
            j += 1
        para = " ".join(l.strip() for l in para_lines)
        if para:
            p = doc.add_paragraph()
            add_inline_runs(p, para)
        i = j

    doc.save(DST)
    print(f"saved: {DST}")
    print(f"size: {DST.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    convert()
